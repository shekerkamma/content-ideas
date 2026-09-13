#!/usr/bin/env python3
"""Build the .pptx that Google Vids imports, optionally split into parts.

Vids caps a Slides import at 45 scenes and silently drops the rest behind a
dialog that looks like a hang, so a longer deck is split here rather than
discovered later. Slides are full-bleed images because Vids imports a deck as
still scenes regardless.
"""
import argparse
import json
import math
import os
import sys


def resolve(cfg, root, slides_dir, s):
    explicit = (cfg.get("slide_files") or {}).get(s)
    if explicit:
        return explicit if os.path.isabs(explicit) else os.path.join(root, explicit)
    key = int(s) if str(s).isdigit() else s
    return os.path.join(slides_dir, cfg["slide_pattern"].format(slide=key))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--name", required=True)
    ap.add_argument("--max-scenes", type=int, default=45)
    ap.add_argument("--split-at", type=int, default=None,
                    help="1-based slide index to start part B (default: midpoint)")
    a = ap.parse_args()

    from pptx import Presentation
    from pptx.util import Inches, Emu
    from PIL import Image

    cfg = json.load(open(a.config))
    root = os.path.dirname(os.path.abspath(a.config))
    rel = lambda p: p if os.path.isabs(p) else os.path.join(root, p)
    slides_dir = rel(cfg["slides_dir"])
    order = [str(s) for s in cfg["slide_order"]]
    narr = json.load(open(rel(cfg["narration"])))

    imgs = [resolve(cfg, root, slides_dir, s) for s in order]
    for s, p in zip(order, imgs):
        if not os.path.exists(p):
            sys.exit(f"BLOCKED: slide image missing for {s}: {p}")

    n = len(order)
    if n <= a.max_scenes:
        parts = [(a.name, 0, n)]
    else:
        cut = a.split_at - 1 if a.split_at else math.ceil(n / 2)
        if cut > a.max_scenes or (n - cut) > a.max_scenes:
            sys.exit(f"BLOCKED: split at {cut} leaves a part over {a.max_scenes}.")
        parts = [(f"{a.name}-A", 0, cut), (f"{a.name}-B", cut, n)]

    w, h = Image.open(imgs[0]).size
    w_in = 13.333
    h_in = round(w_in * h / w, 3)
    os.makedirs(rel(a.out_dir), exist_ok=True)

    manifest = {"parts": []}
    for name, lo, hi in parts:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(w_in), Inches(h_in)
        blank = prs.slide_layouts[6]
        for s, img in zip(order[lo:hi], imgs[lo:hi]):
            sl = prs.slides.add_slide(blank)
            sl.shapes.add_picture(img, Emu(0), Emu(0),
                                  width=prs.slide_width, height=prs.slide_height)
            if narr.get(s):
                sl.notes_slide.notes_text_frame.text = narr[s]
        out = os.path.join(rel(a.out_dir), f"{name}.pptx")
        prs.save(out)
        # The scene->slide map is what the narration injector reads: Vids
        # scene N carries this slide id, so narration never lands on the
        # wrong scene.
        manifest["parts"].append({
            "name": name, "file": os.path.relpath(out, root),
            "scenes": hi - lo, "slides": order[lo:hi],
            "narration": [narr.get(s, "") for s in order[lo:hi]],
        })
        print(f"  {name}: {hi-lo} scenes  {os.path.getsize(out)/1048576:.1f} MB")

    mp = os.path.join(rel(a.out_dir), f"{a.name}-scenes.json")
    json.dump(manifest, open(mp, "w"), indent=1, ensure_ascii=False)
    print(f"  scene map -> {os.path.relpath(mp, root)}")


if __name__ == "__main__":
    main()
